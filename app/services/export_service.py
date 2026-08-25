from sqlalchemy.orm import Session
from datetime import datetime
from typing import Optional
import os
from pathlib import Path

from .base_service import BaseService
from .report_service import ReportService
from .regulatory_service import RegulatoryService
from .stats_service import StatsService, RISK_TYPE_DISPLAY
from ..models import (
    ExportRecord,
    Batch,
    Report,
    ReInspection,
    TestResult,
    Arbitration,
    RegulatoryInspection,
    RegulatoryDifference,
)
from ..schemas import ExportRequest, ExportRecordResponse, CrossBatchRiskTrendQuery
from ..exceptions import ExportError, BatchNotFoundError, ReportNotFoundError
from ..config import settings

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.units import cm
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


class ExportService(BaseService):
    def __init__(self, db: Session):
        super().__init__(db)
        self.report_service = ReportService(db)
        self.regulatory_service = RegulatoryService(db)
        self.stats_service = StatsService(db)
        self.export_dir = settings.EXPORT_DIR

    def _generate_export_no(self) -> str:
        export_no = self.generate_code("EXP")
        while self.db.query(ExportRecord).filter(ExportRecord.export_no == export_no).first():
            export_no = self.generate_code("EXP")
        return export_no

    def _get_export_type(self, batch_id: Optional[int], report_id: Optional[int]) -> str:
        if report_id:
            return "report"
        if batch_id:
            return "batch"
        return "batch"

    def validate_export_request(self, data: ExportRequest) -> None:
        if data.file_format not in ["xlsx", "pdf"]:
            raise ExportError(f"不支持的导出格式: {data.file_format}")

        if data.file_format == "xlsx" and not EXCEL_AVAILABLE:
            raise ExportError("openpyxl 库未安装，无法导出 Excel")

        if data.file_format == "pdf" and not PDF_AVAILABLE:
            raise ExportError("reportlab 库未安装，无法导出 PDF")

        if not data.batch_id and not data.report_id:
            raise ExportError("请指定批次ID或报告ID")

    def export_report(self, data: ExportRequest) -> ExportRecord:
        self.validate_export_request(data)

        export_no = self._generate_export_no()
        export_type = self._get_export_type(data.batch_id, data.report_id)
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S")

        if data.report_id:
            report = self.report_service.get_report(data.report_id)
            batch = report.batch
            file_name = f"质量报告_{report.report_no}_{timestamp}.{data.file_format}"
        else:
            batch = self.db.query(Batch).filter(Batch.id == data.batch_id).first()
            if not batch:
                raise BatchNotFoundError(batch_id=data.batch_id)
            file_name = f"批次质量档案_{batch.batch_no}_{timestamp}.{data.file_format}"

        file_path = os.path.join(self.export_dir, file_name)

        if data.file_format == "xlsx":
            self._export_excel(batch, data, file_path)
        else:
            self._export_pdf(batch, data, file_path)

        file_size = os.path.getsize(file_path) if os.path.exists(file_path) else None

        export_record = ExportRecord(
            export_no=export_no,
            export_type=export_type,
            batch_id=batch.id,
            report_id=data.report_id,
            file_format=data.file_format,
            file_name=file_name,
            file_path=file_path,
            file_size=file_size,
            exported_by=data.exported_by,
            exported_at=datetime.now(),
            include_original=data.include_original,
            include_reinspection=data.include_reinspection,
            include_final=data.include_final,
            include_arbitration=data.include_arbitration,
            include_regulatory=data.include_regulatory,
        )

        self.db.add(export_record)
        self.db.commit()
        self.db.refresh(export_record)
        return export_record

    def _get_batch_data(self, batch: Batch, data: ExportRequest) -> dict:
        original_report = None
        reinspection_reports = []
        final_report = None
        arbitration_reports = []
        reinspections = []
        arbitrations = []
        regulatory_inspections = []
        regulatory_differences = []
        risk_summary = None

        if data.include_original:
            original_report = (
                self.db.query(Report)
                .filter(Report.batch_id == batch.id, Report.report_type == "original")
                .first()
            )

        if data.include_reinspection:
            reinspection_reports = (
                self.db.query(Report)
                .filter(Report.batch_id == batch.id, Report.report_type == "reinspection")
                .order_by(Report.created_at)
                .all()
            )
            reinspections = (
                self.db.query(ReInspection)
                .filter(ReInspection.original_sample_id.in_([s.id for s in batch.samples]))
                .order_by(ReInspection.created_at)
                .all()
            )

        if data.include_final:
            final_report = (
                self.db.query(Report)
                .filter(Report.batch_id == batch.id, Report.report_type == "final")
                .first()
            )

        if data.include_arbitration:
            arbitration_reports = (
                self.db.query(Report)
                .filter(Report.batch_id == batch.id, Report.report_type == "arbitration")
                .order_by(Report.created_at)
                .all()
            )
            arbitrations = (
                self.db.query(Arbitration)
                .filter(Arbitration.batch_id == batch.id)
                .order_by(Arbitration.created_at)
                .all()
            )

        if data.include_regulatory:
            regulatory_inspections = (
                self.db.query(RegulatoryInspection)
                .filter(RegulatoryInspection.batch_id == batch.id)
                .order_by(RegulatoryInspection.created_at)
                .all()
            )
            regulatory_differences = (
                self.db.query(RegulatoryDifference)
                .filter(RegulatoryDifference.batch_id == batch.id)
                .order_by(RegulatoryDifference.created_at)
                .all()
            )
            risk_summary = self.stats_service.analyze_cross_batch_risks(
                CrossBatchRiskTrendQuery(
                    scope_dimension="product",
                    scope_value=batch.product_name,
                )
            )

        return {
            "batch": batch,
            "original_report": original_report,
            "reinspection_reports": reinspection_reports,
            "final_report": final_report,
            "arbitration_reports": arbitration_reports,
            "reinspections": reinspections,
            "arbitrations": arbitrations,
            "regulatory_inspections": regulatory_inspections,
            "regulatory_differences": regulatory_differences,
            "risk_summary": risk_summary,
        }

    def _export_excel(self, batch: Batch, data: ExportRequest, file_path: str) -> None:
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "质量报告"

            header_font = Font(bold=True, size=14, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            subheader_font = Font(bold=True, size=11)
            center_align = Alignment(horizontal="center", vertical="center")

            row = 1
            ws.cell(row=row, column=1, value="质检样品留样与复检流程服务 - 质量报告")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            ws.cell(row=row, column=1).font = Font(bold=True, size=16)
            ws.cell(row=row, column=1).alignment = center_align
            row += 2

            ws.cell(row=row, column=1, value="批次信息")
            ws.cell(row=row, column=1).font = header_font
            ws.cell(row=row, column=1).fill = header_fill
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
            row += 1

            risk_level_map = {"low": "低", "medium": "中", "high": "高"}
            risk_level_display = risk_level_map.get(batch.risk_level, batch.risk_level)
            batch_info = [
                ["批次号", batch.batch_no, "产品名称", batch.product_name],
                ["生产日期", batch.production_date.strftime("%Y-%m-%d"), "生产数量", str(batch.quantity)],
                ["生产线", batch.production_line or "-", "批次状态", batch.status],
                ["最终结果", batch.final_result or "-", "是否放行", "是" if batch.released else "否"],
                ["风险等级", risk_level_display, "风险评分", str(batch.risk_score)],
                ["是否冻结", "是" if batch.frozen else "否", "冻结原因", batch.frozen_reason or "-"],
                ["是否撤销放行", "是" if batch.release_revoked else "否", "撤销原因", batch.release_revoked_reason or "-"],
            ]
            for info_row in batch_info:
                for col, val in enumerate(info_row, 1):
                    ws.cell(row=row, column=col, value=val)
                    if col % 2 == 1:
                        ws.cell(row=row, column=col).font = subheader_font
                row += 1
            row += 1

            batch_data = self._get_batch_data(batch, data)

            if data.include_original and batch_data["original_report"]:
                report = batch_data["original_report"]
                row = self._write_report_to_excel(ws, row, "原始检测报告", report, header_font, header_fill, subheader_font)

            if data.include_reinspection:
                for report in batch_data["reinspection_reports"]:
                    row = self._write_report_to_excel(ws, row, "复检检测报告", report, header_font, header_fill, subheader_font)
                row += 1

                if batch_data["reinspections"]:
                    ws.cell(row=row, column=1, value="复检记录")
                    ws.cell(row=row, column=1).font = header_font
                    ws.cell(row=row, column=1).fill = header_fill
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                    row += 1

                    headers = ["复检编号", "原因", "申请人", "申请时间", "状态", "是否有差异", "最终判定"]
                    for col, header in enumerate(headers, 1):
                        ws.cell(row=row, column=col, value=header)
                        ws.cell(row=row, column=col).font = subheader_font
                    row += 1

                    for ri in batch_data["reinspections"]:
                        ws.cell(row=row, column=1, value=ri.reinspection_code)
                        ws.cell(row=row, column=2, value=ri.reason)
                        ws.cell(row=row, column=3, value=ri.applicant)
                        ws.cell(row=row, column=4, value=ri.apply_time.strftime("%Y-%m-%d %H:%M:%S"))
                        ws.cell(row=row, column=5, value=ri.status)
                        ws.cell(row=row, column=6, value="是" if ri.difference_identified else "否" if ri.difference_identified is not None else "-")
                        ws.cell(row=row, column=7, value=ri.final_judgment or "-")
                        row += 1
                    row += 1

            if data.include_final and batch_data["final_report"]:
                report = batch_data["final_report"]
                row = self._write_report_to_excel(ws, row, "最终报告", report, header_font, header_fill, subheader_font)

            if data.include_arbitration:
                for report in batch_data["arbitration_reports"]:
                    row = self._write_report_to_excel(ws, row, "仲裁报告", report, header_font, header_fill, subheader_font)
                row += 1

                if batch_data["arbitrations"]:
                    ws.cell(row=row, column=1, value="仲裁记录")
                    ws.cell(row=row, column=1).font = header_font
                    ws.cell(row=row, column=1).fill = header_fill
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                    row += 1

                    headers = ["仲裁编号", "争议来源", "仲裁负责人", "状态", "处理意见", "最终判定", "创建时间", "完成时间"]
                    for col, header in enumerate(headers, 1):
                        ws.cell(row=row, column=col, value=header)
                        ws.cell(row=row, column=col).font = subheader_font
                    row += 1

                    for arb in batch_data["arbitrations"]:
                        ws.cell(row=row, column=1, value=arb.arbitration_code)
                        ws.cell(row=row, column=2, value=arb.dispute_source)
                        ws.cell(row=row, column=3, value=arb.arbitrator)
                        ws.cell(row=row, column=4, value=arb.status)
                        ws.cell(row=row, column=5, value=arb.handling_opinion or "-")
                        ws.cell(row=row, column=6, value=arb.final_judgment or "-")
                        ws.cell(row=row, column=7, value=arb.created_at.strftime("%Y-%m-%d %H:%M:%S"))
                        ws.cell(row=row, column=8, value=arb.completed_at.strftime("%Y-%m-%d %H:%M:%S") if arb.completed_at else "-")
                        row += 1
                    row += 1

            if data.include_regulatory:
                if batch_data["regulatory_inspections"]:
                    ws.cell(row=row, column=1, value="监管抽检记录")
                    ws.cell(row=row, column=1).font = header_font
                    ws.cell(row=row, column=1).fill = header_fill
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                    row += 1

                    headers = ["抽检编号", "抽检机构", "抽检人员", "抽检日期", "监管抽检结论", "企业原判定", "状态", "是否生成差异"]
                    for col, header in enumerate(headers, 1):
                        ws.cell(row=row, column=col, value=header)
                        ws.cell(row=row, column=col).font = subheader_font
                    row += 1

                    for ri in batch_data["regulatory_inspections"]:
                        ri_resp = self.regulatory_service.to_inspection_response(ri)
                        ws.cell(row=row, column=1, value=ri.inspection_code)
                        ws.cell(row=row, column=2, value=ri.inspection_agency)
                        ws.cell(row=row, column=3, value=ri.inspector)
                        ws.cell(row=row, column=4, value=ri.inspection_date.strftime("%Y-%m-%d"))
                        ws.cell(row=row, column=5, value=ri.overall_conclusion or "-")
                        ws.cell(row=row, column=6, value=ri_resp.enterprise_judgment if hasattr(ri_resp, 'enterprise_judgment') else (batch.final_result or "-"))
                        ws.cell(row=row, column=7, value=ri.status)
                        ws.cell(row=row, column=8, value="是" if ri.difference_generated else "否")
                        row += 1

                        if ri.item_results:
                            ws.cell(row=row, column=2, value="检测项明细")
                            ws.cell(row=row, column=2).font = subheader_font
                            row += 1
                            item_headers = ["", "检测项名称", "企业检测值", "企业判定", "监管检测值", "监管判定", "是否一致", "差异说明"]
                            for col, header in enumerate(item_headers, 1):
                                ws.cell(row=row, column=col, value=header)
                                ws.cell(row=row, column=col).font = subheader_font
                            row += 1
                            for idx, item in enumerate(ri.item_results):
                                ti = item.test_item
                                ws.cell(row=row, column=1, value=f"[{idx+1}]")
                                ws.cell(row=row, column=2, value=ti.item_name if ti else f"ID:{item.test_item_id}")
                                ws.cell(row=row, column=3, value=item.enterprise_value or "-")
                                ws.cell(row=row, column=4, value=item.enterprise_judgment or "-")
                                ws.cell(row=row, column=5, value=item.regulatory_value)
                                ws.cell(row=row, column=6, value=item.regulatory_judgment or "-")
                                ws.cell(row=row, column=7, value="是" if item.is_consistent else "否" if item.is_consistent is not None else "-")
                                ws.cell(row=row, column=8, value=item.difference_detail or "-")
                                if item.is_consistent == False:
                                    for col in range(1, 9):
                                        ws.cell(row=row, column=col).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                                row += 1
                    row += 1

                if batch_data["regulatory_differences"]:
                    ws.cell(row=row, column=1, value="监管差异单")
                    ws.cell(row=row, column=1).font = header_font
                    ws.cell(row=row, column=1).fill = header_fill
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                    row += 1

                    diff_headers = ["差异编号", "差异类型", "企业判定", "监管判定", "差异原因", "处置措施", "处置状态", "批次冻结/撤销放行"]
                    for col, header in enumerate(diff_headers, 1):
                        ws.cell(row=row, column=col, value=header)
                        ws.cell(row=row, column=col).font = subheader_font
                    row += 1

                    for rd in batch_data["regulatory_differences"]:
                        ws.cell(row=row, column=1, value=rd.difference_code)
                        ws.cell(row=row, column=2, value=rd.difference_type or "-")
                        ws.cell(row=row, column=3, value=rd.enterprise_judgment or "-")
                        ws.cell(row=row, column=4, value=rd.regulatory_judgment or "-")
                        ws.cell(row=row, column=5, value=rd.difference_reason or "-")
                        ws.cell(row=row, column=6, value=rd.handling_measures or "-")
                        ws.cell(row=row, column=7, value=rd.handling_status)
                        ws.cell(row=row, column=8, value=f"冻结:{'是' if rd.batch_frozen else '否'};撤销:{'是' if rd.release_revoked else '否'}")
                        row += 1
                    row += 1

                if batch_data["risk_summary"] and batch_data["risk_summary"].details:
                    ws.cell(row=row, column=1, value="跨批次风险摘要")
                    ws.cell(row=row, column=1).font = header_font
                    ws.cell(row=row, column=1).fill = header_fill
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
                    row += 1

                    if batch_data["risk_summary"].summary:
                        ws.cell(row=row, column=1, value="风险汇总")
                        ws.cell(row=row, column=1).font = subheader_font
                        row += 1
                        s_headers = ["风险类型", "中文名称", "影响批次数", "高风险", "中风险", "低风险"]
                        for col, header in enumerate(s_headers, 1):
                            ws.cell(row=row, column=col, value=header)
                            ws.cell(row=row, column=col).font = subheader_font
                        row += 1
                        for s in batch_data["risk_summary"].summary:
                            ws.cell(row=row, column=1, value=s.risk_type)
                            ws.cell(row=row, column=2, value=s.risk_type_display)
                            ws.cell(row=row, column=3, value=s.affected_count)
                            ws.cell(row=row, column=4, value=s.high_risk_count)
                            ws.cell(row=row, column=5, value=s.medium_risk_count)
                            ws.cell(row=row, column=6, value=s.low_risk_count)
                            row += 1
                        row += 1

                    ws.cell(row=row, column=1, value="风险明细")
                    ws.cell(row=row, column=1).font = subheader_font
                    row += 1
                    d_headers = ["风险类型", "维度", "维度值", "风险等级", "影响批次数", "描述"]
                    for col, header in enumerate(d_headers, 1):
                        ws.cell(row=row, column=col, value=header)
                        ws.cell(row=row, column=col).font = subheader_font
                    row += 1
                    for d in batch_data["risk_summary"].details:
                        ws.cell(row=row, column=1, value=d.risk_type_display)
                        ws.cell(row=row, column=2, value=d.scope_dimension or "-")
                        ws.cell(row=row, column=3, value=d.scope_value or "-")
                        risk_level_map_xlsx = {"high": "高", "medium": "中", "low": "低"}
                        ws.cell(row=row, column=4, value=risk_level_map_xlsx.get(d.risk_level, d.risk_level))
                        ws.cell(row=row, column=5, value=d.affected_batch_count)
                        ws.cell(row=row, column=6, value=d.description or "-")
                        if d.risk_level == "high":
                            for col in range(1, 7):
                                ws.cell(row=row, column=col).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                        elif d.risk_level == "medium":
                            for col in range(1, 7):
                                ws.cell(row=row, column=col).fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
                        row += 1
                    row += 1

            for col in range(1, 9):
                ws.column_dimensions[get_column_letter(col)].width = 22

            wb.save(file_path)
        except Exception as e:
            raise ExportError(f"Excel 导出失败: {str(e)}")

    def _write_report_to_excel(self, ws, row, title, report, header_font, header_fill, subheader_font):
        ws.cell(row=row, column=1, value=f"{title} - {report.report_no}")
        ws.cell(row=row, column=1).font = header_font
        ws.cell(row=row, column=1).fill = header_fill
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1

        ws.cell(row=row, column=1, value="报告类型")
        ws.cell(row=row, column=2, value=report.report_type)
        ws.cell(row=row, column=3, value="总判定")
        ws.cell(row=row, column=4, value=report.overall_judgment)
        ws.cell(row=row, column=5, value="出具人")
        ws.cell(row=row, column=6, value=report.issued_by or "-")
        ws.cell(row=row, column=7, value="出具时间")
        ws.cell(row=row, column=8, value=report.issued_at.strftime("%Y-%m-%d %H:%M:%S"))
        for col in range(1, 9):
            if col % 2 == 1:
                ws.cell(row=row, column=col).font = subheader_font
        row += 1
        row += 1

        ws.cell(row=row, column=1, value="检测项目明细")
        ws.cell(row=row, column=1).font = subheader_font
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1

        headers = ["检测项编码", "检测项名称", "检测值", "标准值", "下限", "上限", "单位", "判定"]
        for col, header in enumerate(headers, 1):
            ws.cell(row=row, column=col, value=header)
            ws.cell(row=row, column=col).font = subheader_font
        row += 1

        report_resp = self.report_service.to_response(report)
        for tr in report_resp.test_results:
            ws.cell(row=row, column=1, value=tr.test_item_code)
            ws.cell(row=row, column=2, value=tr.test_item_name)
            ws.cell(row=row, column=3, value=tr.test_value)
            ws.cell(row=row, column=4, value=tr.standard_value or "-")
            ws.cell(row=row, column=5, value=str(tr.lower_limit) if tr.lower_limit is not None else "-")
            ws.cell(row=row, column=6, value=str(tr.upper_limit) if tr.upper_limit is not None else "-")
            ws.cell(row=row, column=7, value=tr.unit or "-")
            ws.cell(row=row, column=8, value=tr.judgment or "-")
            if tr.judgment == "fail":
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            row += 1
        row += 2

        return row

    def _export_pdf(self, batch: Batch, data: ExportRequest, file_path: str) -> None:
        try:
            doc = SimpleDocTemplate(file_path, pagesize=landscape(A4), leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
            story = []
            styles = getSampleStyleSheet()

            title_style = ParagraphStyle(
                "CustomTitle",
                parent=styles["Title"],
                fontSize=18,
                textColor=colors.HexColor("#4472C4"),
                spaceAfter=20,
            )
            heading_style = ParagraphStyle(
                "CustomHeading",
                parent=styles["Heading2"],
                fontSize=14,
                textColor=colors.white,
                backColor=colors.HexColor("#4472C4"),
                padding=8,
                spaceAfter=12,
            )
            subheading_style = ParagraphStyle(
                "CustomSubHeading",
                parent=styles["Heading3"],
                fontSize=12,
                textColor=colors.HexColor("#4472C4"),
                spaceAfter=8,
            )
            normal_style = styles["Normal"]

            story.append(Paragraph("质检样品留样与复检流程服务 - 质量报告", title_style))
            story.append(Spacer(1, 0.5*cm))

            story.append(Paragraph("批次信息", heading_style))

            risk_level_map = {"low": "低风险", "medium": "中风险", "high": "高风险"}
            risk_level_display = risk_level_map.get(batch.risk_level, batch.risk_level)
            batch_data = [
                ["批次号", batch.batch_no, "产品名称", batch.product_name],
                ["生产日期", batch.production_date.strftime("%Y-%m-%d"), "生产数量", str(batch.quantity)],
                ["生产线", batch.production_line or "-", "批次状态", batch.status],
                ["最终结果", batch.final_result or "-", "是否放行", "是" if batch.released else "否"],
                ["风险等级", risk_level_display, "风险评分", str(batch.risk_score)],
                ["是否冻结", "是" if batch.frozen else "否", "冻结原因", batch.frozen_reason or "-"],
                ["是否撤销放行", "是" if batch.release_revoked else "否", "撤销原因", batch.release_revoked_reason or "-"],
            ]
            batch_table = Table(batch_data, colWidths=[3*cm, 5*cm, 3*cm, 5*cm])
            batch_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#D9E2F3")),
                ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#D9E2F3")),
                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(batch_table)
            story.append(Spacer(1, 0.8*cm))

            batch_data_dict = self._get_batch_data(batch, data)

            if data.include_original and batch_data_dict["original_report"]:
                story = self._write_report_to_pdf(story, "原始检测报告", batch_data_dict["original_report"], heading_style, subheading_style, normal_style)

            if data.include_reinspection:
                for report in batch_data_dict["reinspection_reports"]:
                    story = self._write_report_to_pdf(story, "复检检测报告", report, heading_style, subheading_style, normal_style)

                if batch_data_dict["reinspections"]:
                    story.append(Paragraph("复检记录", heading_style))
                    reinspection_data = [["复检编号", "原因", "申请人", "申请时间", "状态", "差异", "最终判定"]]
                    for ri in batch_data_dict["reinspections"]:
                        reinspection_data.append([
                            ri.reinspection_code,
                            ri.reason,
                            ri.applicant,
                            ri.apply_time.strftime("%Y-%m-%d %H:%M"),
                            ri.status,
                            "是" if ri.difference_identified else "否" if ri.difference_identified is not None else "-",
                            ri.final_judgment or "-",
                        ])
                    reinspection_table = Table(reinspection_data, colWidths=[3.2*cm, 4*cm, 2*cm, 2.8*cm, 1.8*cm, 1.5*cm, 2*cm])
                    reinspection_table.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("PADDING", (0, 0), (-1, -1), 4),
                    ]))
                    story.append(reinspection_table)
                    story.append(Spacer(1, 0.8*cm))

            if data.include_final and batch_data_dict["final_report"]:
                story = self._write_report_to_pdf(story, "最终报告", batch_data_dict["final_report"], heading_style, subheading_style, normal_style)

            if data.include_arbitration:
                for report in batch_data_dict["arbitration_reports"]:
                    story = self._write_report_to_pdf(story, "仲裁报告", report, heading_style, subheading_style, normal_style)

                if batch_data_dict["arbitrations"]:
                    story.append(Paragraph("仲裁记录", heading_style))
                    arbitration_data = [["仲裁编号", "争议来源", "仲裁负责人", "状态", "最终判定"]]
                    for arb in batch_data_dict["arbitrations"]:
                        arbitration_data.append([
                            arb.arbitration_code,
                            arb.dispute_source[:40] + "..." if len(arb.dispute_source) > 40 else arb.dispute_source,
                            arb.arbitrator,
                            arb.status,
                            arb.final_judgment or "-",
                        ])
                    arbitration_table = Table(arbitration_data, colWidths=[3.5*cm, 5*cm, 2.5*cm, 2*cm, 2*cm])
                    arbitration_table.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("PADDING", (0, 0), (-1, -1), 4),
                    ]))
                    story.append(arbitration_table)
                    story.append(Spacer(1, 0.8*cm))

            if data.include_regulatory:
                if batch_data_dict["regulatory_inspections"]:
                    story.append(Paragraph("监管抽检记录", heading_style))
                    for ri in batch_data_dict["regulatory_inspections"]:
                        ri_resp = self.regulatory_service.to_inspection_response(ri)
                        story.append(Paragraph(f"抽检编号：{ri.inspection_code}", subheading_style))
                        insp_data = [
                            ["抽检机构", ri.inspection_agency, "抽检人员", ri.inspector],
                            ["抽检日期", ri.inspection_date.strftime("%Y-%m-%d"), "状态", ri.status],
                            ["监管抽检结论", ri.overall_conclusion or "-",
                             "企业原判定", ri_resp.enterprise_judgment if hasattr(ri_resp, 'enterprise_judgment') else (batch.final_result or "-")],
                            ["是否生成差异", "是" if ri.difference_generated else "否", "", ""],
                        ]
                        insp_table = Table(insp_data, colWidths=[3*cm, 5*cm, 3*cm, 5*cm])
                        insp_table.setStyle(TableStyle([
                            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#D9E2F3")),
                            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#D9E2F3")),
                            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                            ("FONTSIZE", (0, 0), (-1, -1), 9),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("PADDING", (0, 0), (-1, -1), 4),
                        ]))
                        story.append(insp_table)
                        story.append(Spacer(1, 0.3*cm))

                        if ri.item_results:
                            story.append(Paragraph("检测项明细", subheading_style))
                            item_data = [["序号", "检测项名称", "企业值", "企业判定", "监管值", "监管判定", "一致", "差异说明"]]
                            for idx, item in enumerate(ri.item_results):
                                ti = item.test_item
                                item_data.append([
                                    str(idx+1),
                                    ti.item_name if ti else f"ID:{item.test_item_id}",
                                    item.enterprise_value or "-",
                                    item.enterprise_judgment or "-",
                                    item.regulatory_value,
                                    item.regulatory_judgment or "-",
                                    "是" if item.is_consistent else "否" if item.is_consistent is not None else "-",
                                    (item.difference_detail or "")[:30],
                                ])
                            item_table = Table(item_data, colWidths=[1.2*cm, 3.5*cm, 2*cm, 2*cm, 2*cm, 2*cm, 1.3*cm, 3*cm])
                            style_cmds = [
                                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                                ("FONTSIZE", (0, 0), (-1, -1), 8),
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("PADDING", (0, 0), (-1, -1), 3),
                            ]
                            for ridx, item in enumerate(ri.item_results):
                                if item.is_consistent == False:
                                    style_cmds.append(("BACKGROUND", (0, ridx+1), (-1, ridx+1), colors.HexColor("#FFC7CE")))
                            item_table.setStyle(TableStyle(style_cmds))
                            story.append(item_table)
                        story.append(Spacer(1, 0.6*cm))

                if batch_data_dict["regulatory_differences"]:
                    story.append(Paragraph("监管差异单", heading_style))
                    diff_data = [["差异编号", "企业判定", "监管判定", "差异原因", "处置措施", "处理状态", "冻结/撤销"]]
                    for rd in batch_data_dict["regulatory_differences"]:
                        diff_data.append([
                            rd.difference_code,
                            rd.enterprise_judgment or "-",
                            rd.regulatory_judgment or "-",
                            (rd.difference_reason or "")[:30],
                            (rd.handling_measures or "")[:30],
                            rd.handling_status,
                            f"冻:{'是' if rd.batch_frozen else '否'}\n撤:{'是' if rd.release_revoked else '否'}",
                        ])
                    diff_table = Table(diff_data, colWidths=[2.8*cm, 2.2*cm, 2.2*cm, 3.2*cm, 3.2*cm, 2*cm, 2.2*cm])
                    diff_table.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("PADDING", (0, 0), (-1, -1), 4),
                    ]))
                    story.append(diff_table)
                    story.append(Spacer(1, 0.8*cm))

                if batch_data_dict["risk_summary"] and batch_data_dict["risk_summary"].details:
                    story.append(Paragraph("跨批次风险摘要", heading_style))
                    if batch_data_dict["risk_summary"].summary:
                        story.append(Paragraph("风险汇总", subheading_style))
                        s_data = [["风险类型", "中文名称", "影响批次数", "高风险", "中风险", "低风险"]]
                        for s in batch_data_dict["risk_summary"].summary:
                            s_data.append([
                                s.risk_type,
                                s.risk_type_display,
                                str(s.affected_count),
                                str(s.high_risk_count),
                                str(s.medium_risk_count),
                                str(s.low_risk_count),
                            ])
                        s_table = Table(s_data, colWidths=[3*cm, 3*cm, 2.5*cm, 2*cm, 2*cm, 2*cm])
                        s_table.setStyle(TableStyle([
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                            ("FONTSIZE", (0, 0), (-1, -1), 9),
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                            ("PADDING", (0, 0), (-1, -1), 4),
                        ]))
                        story.append(s_table)
                        story.append(Spacer(1, 0.4*cm))

                    story.append(Paragraph("风险明细", subheading_style))
                    d_data = [["风险类型", "维度", "维度值", "风险等级", "批次数", "描述"]]
                    style_cmds_d = [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("PADDING", (0, 0), (-1, -1), 3),
                    ]
                    for ridx, d in enumerate(batch_data_dict["risk_summary"].details):
                        risk_level_disp = {"high": "高", "medium": "中", "low": "低"}.get(d.risk_level, d.risk_level)
                        d_data.append([
                            d.risk_type_display,
                            d.scope_dimension or "-",
                            (d.scope_value or "")[:15],
                            risk_level_disp,
                            str(d.affected_batch_count),
                            (d.description or "")[:45],
                        ])
                        if d.risk_level == "high":
                            style_cmds_d.append(("BACKGROUND", (0, ridx+1), (-1, ridx+1), colors.HexColor("#FFC7CE")))
                        elif d.risk_level == "medium":
                            style_cmds_d.append(("BACKGROUND", (0, ridx+1), (-1, ridx+1), colors.HexColor("#FFEB9C")))
                    d_table = Table(d_data, colWidths=[2.8*cm, 2*cm, 2.5*cm, 2*cm, 1.5*cm, 4.2*cm])
                    d_table.setStyle(TableStyle(style_cmds_d))
                    story.append(d_table)
                    story.append(Spacer(1, 0.8*cm))

            doc.build(story)
        except Exception as e:
            raise ExportError(f"PDF 导出失败: {str(e)}")

    def _write_report_to_pdf(self, story, title, report, heading_style, subheading_style, normal_style):
        story.append(Paragraph(f"{title} - {report.report_no}", heading_style))

        report_info = [
            ["报告类型", report.report_type, "总判定", report.overall_judgment],
            ["出具人", report.issued_by or "-", "出具时间", report.issued_at.strftime("%Y-%m-%d %H:%M:%S")],
        ]
        report_table = Table(report_info, colWidths=[2.5*cm, 5*cm, 2.5*cm, 5*cm])
        report_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#D9E2F3")),
            ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#D9E2F3")),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 10),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("PADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(report_table)
        story.append(Spacer(1, 0.5*cm))

        story.append(Paragraph("检测项目明细", subheading_style))

        report_resp = self.report_service.to_response(report)
        test_data = [["检测项编码", "检测项名称", "检测值", "标准值", "下限", "上限", "单位", "判定"]]
        for tr in report_resp.test_results:
            test_data.append([
                tr.test_item_code,
                tr.test_item_name,
                tr.test_value,
                tr.standard_value or "-",
                str(tr.lower_limit) if tr.lower_limit is not None else "-",
                str(tr.upper_limit) if tr.upper_limit is not None else "-",
                tr.unit or "-",
                tr.judgment or "-",
            ])
        test_table = Table(test_data, colWidths=[2*cm, 3*cm, 2*cm, 2.5*cm, 1.5*cm, 1.5*cm, 1.5*cm, 1.5*cm])
        style = TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("PADDING", (0, 0), (-1, -1), 4),
        ])
        for i, tr in enumerate(report_resp.test_results, 1):
            if tr.judgment == "fail":
                style.add("BACKGROUND", (0, i), (-1, i), colors.HexColor("#FFC7CE"))
        test_table.setStyle(style)
        story.append(test_table)
        story.append(Spacer(1, 1*cm))

        return story

    def list_export_records(self, skip: int = 0, limit: int = 100, batch_id: Optional[int] = None) -> list:
        query = self.db.query(ExportRecord)
        if batch_id:
            query = query.filter(ExportRecord.batch_id == batch_id)
        return query.order_by(ExportRecord.created_at.desc()).offset(skip).limit(limit).all()

    def to_response(self, record: ExportRecord) -> ExportRecordResponse:
        return ExportRecordResponse(
            id=record.id,
            export_no=record.export_no,
            export_type=record.export_type,
            batch_id=record.batch_id,
            batch_no=record.batch.batch_no if record.batch else None,
            report_id=record.report_id,
            report_no=record.report.report_no if record.report else None,
            file_format=record.file_format,
            file_name=record.file_name,
            file_path=record.file_path,
            file_size=record.file_size,
            exported_by=record.exported_by,
            exported_at=record.exported_at,
            include_original=record.include_original,
            include_reinspection=record.include_reinspection,
            include_final=record.include_final,
            include_arbitration=record.include_arbitration,
            include_regulatory=record.include_regulatory,
            download_url=f"/exports/download/{record.export_no}",
        )
